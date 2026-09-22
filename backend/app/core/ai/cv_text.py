"""Extrakcia čistého textu zo životopisu (PDF alebo DOCX).

Text posielame modelu namiesto samotného súboru: je to lacnejšie a funguje
aj pre DOCX. Skenované PDF bez textovej vrstvy vrátia prázdny text; v tom
prípade sa uchádzač hodnotí len z chatu a zdôvodnenie to uvedie.
"""

from __future__ import annotations

import io
import logging
import re

logger = logging.getLogger(__name__)

MAX_CHARS = 15_000
TRUNCATED_MARKER = "\n[… text životopisu bol skrátený …]"
SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


class UnsupportedCvFormat(ValueError):
    pass


def extract_text(data: bytes, filename: str) -> str:
    """Vráť text zo súboru podľa prípony. Prázdny reťazec, ak sa text nepodarilo získať."""
    ext = _extension(filename)
    if ext == ".pdf":
        text = _from_pdf(data)
    elif ext == ".docx":
        text = _from_docx(data)
    else:
        raise UnsupportedCvFormat(f"Nepodporovaný formát životopisu: {filename}")
    return _normalize(text)


def _extension(filename: str) -> str:
    dot = filename.rfind(".")
    return filename[dot:].lower() if dot >= 0 else ""


def _from_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:  # noqa: BLE001
                logger.warning("CV PDF je zašifrované, text sa nedá prečítať")
                return ""
    except Exception:  # noqa: BLE001
        logger.warning("CV PDF sa nepodarilo otvoriť", exc_info=True)
        return ""

    pages: list[str] = []
    for index, page in enumerate(reader.pages):
        try:
            pages.append(page.extract_text() or "")
        except Exception:  # noqa: BLE001
            logger.warning("Strana %d CV PDF sa nedala prečítať", index + 1, exc_info=True)
        if sum(len(p) for p in pages) > MAX_CHARS:
            break
    return "\n\n".join(pages)


def _from_docx(data: bytes) -> str:
    from docx import Document

    try:
        document = Document(io.BytesIO(data))
    except Exception:  # noqa: BLE001
        logger.warning("CV DOCX sa nepodarilo otvoriť", exc_info=True)
        return ""

    lines: list[str] = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            lines.append(" | ".join(c for c in cells if c))
    return "\n".join(lines)


def _normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) > MAX_CHARS:
        text = text[:MAX_CHARS].rstrip() + TRUNCATED_MARKER
    return text
