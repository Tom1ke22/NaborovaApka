"""Testy extrakcie textu z PDF a DOCX životopisov."""

import io

import pytest
from docx import Document

from app.core.ai.cv_text import MAX_CHARS, TRUNCATED_MARKER, UnsupportedCvFormat, extract_text


def build_pdf(text: str) -> bytes:
    """Minimálne platné PDF s jednou stranou a jedným textovým objektom."""
    content = f"BT /F1 12 Tf 20 100 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 200] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref_pos = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode() + b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_pos}\n%%EOF\n"
    ).encode()
    return bytes(out)


def build_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    document = Document()
    for p in paragraphs:
        document.add_paragraph(p)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_pdf_text_is_extracted():
    text = extract_text(build_pdf("Kuchar 3 roky praxe"), "cv.PDF")
    assert "Kuchar 3 roky praxe" in text


def test_docx_paragraphs_and_tables_are_extracted():
    data = build_docx(
        ["Ján Novák", "Kuchár, 3 roky praxe"],
        table=[["Vzdelanie", "SOŠ hotelová"], ["Jazyky", "angličtina B1"]],
    )
    text = extract_text(data, "cv.docx")
    assert "Ján Novák" in text
    assert "Kuchár, 3 roky praxe" in text
    assert "Vzdelanie | SOŠ hotelová" in text
    assert "angličtina B1" in text


def test_unsupported_extension_raises():
    with pytest.raises(UnsupportedCvFormat):
        extract_text(b"whatever", "cv.txt")


def test_corrupt_files_return_empty_text_instead_of_raising():
    assert extract_text(b"this is not a pdf", "cv.pdf") == ""
    assert extract_text(b"this is not a docx", "cv.docx") == ""


def test_long_text_is_truncated_with_marker():
    data = build_docx(["x" * 1000] * 30)
    text = extract_text(data, "cv.docx")
    assert text.endswith(TRUNCATED_MARKER)
    assert len(text) <= MAX_CHARS + len(TRUNCATED_MARKER)


def test_whitespace_is_normalized():
    data = build_docx(["Meno:   Ján  ", "", "", "", "Prax:\t3 roky"])
    text = extract_text(data, "cv.docx")
    assert text == "Meno: Ján\n\nPrax: 3 roky"
